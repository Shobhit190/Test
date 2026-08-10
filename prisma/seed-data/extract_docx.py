import docx
import json
import os
import sys

# Usage: python3 extract_docx.py <Competency_Dictionary_*.docx>
# Regenerates competency-dictionary.json from the source Word file.
SRC = sys.argv[1] if len(sys.argv) > 1 else "Competency_Dictionary_24_Mar_26.docx"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "competency-dictionary.json")

d = docx.Document(SRC)

# --- Table 0: S.No / Cluster / Sub-Cluster / Competency list ---
listing = []
for row in d.tables[0].rows[1:]:
    sno, cluster, subcluster, name = [c.text.strip() for c in row.cells]
    listing.append(
        {
            "sno": sno,
            "cluster": cluster,
            "subCluster": subcluster,
            "name": name.replace(" -", "-").replace("  ", " ").strip(),
        }
    )

core_entries = [e for e in listing if e["subCluster"] == "Core"]
numbered_entries = [e for e in listing if e["sno"].isdigit()]
numbered_entries.sort(key=lambda e: int(e["sno"]))
assert len(numbered_entries) == 19, len(numbered_entries)

LEVEL_CODES = ["A", "B", "C", "D", "E"]


def normalize(name: str) -> str:
    return " ".join(name.replace(" -", "-").split()).strip()


core_names_normalized = {normalize(e["name"]) for e in core_entries}
# Known aliases between the "Core" listing names and the cluster listing names
alias_map = {
    "Student/Education-First Integrity": "Student-First Integrity",
}

# --- Table 1: Proficiency scale ---
scale = []
for row in d.tables[1].rows[1:]:
    level, sublevel, levelname, desc = [c.text.strip() for c in row.cells]
    scale.append(
        {"level": level, "subLevel": int(sublevel), "levelName": levelname, "description": desc}
    )

# --- Competency detail tables: shape 7x3, in the same order as numbered_entries ---
detail_tables = [t for t in d.tables if len(t.rows) == 7 and len(t.columns) == 3]
assert len(detail_tables) == 19, len(detail_tables)

import re

DEF_LABEL_RE = re.compile(r"^definition:?\s*(.*)$", re.IGNORECASE)
KB_LABEL_RE = re.compile(r"^(key behaviours|behavioural indicators|behaviour indicators):?\s*(.*)$", re.IGNORECASE)
NAME_PREFIX_RE = re.compile(r"^(?:[ivx]+\.|\d+\.)\s+", re.IGNORECASE)

competencies = []
for entry, table in zip(numbered_entries, detail_tables):
    header_cell = table.rows[0].cells[0]
    paras = [p.text.strip() for p in header_cell.paragraphs if p.text.strip()]
    name = NAME_PREFIX_RE.sub("", paras[0]).strip()

    def_idx = None
    def_inline = ""
    for i, p in enumerate(paras[1:], start=1):
        m = DEF_LABEL_RE.match(p)
        if m:
            def_idx = i
            def_inline = m.group(1).strip()
            break

    kb_idx = None
    kb_inline = ""
    for i, p in enumerate(paras[1:], start=1):
        m = KB_LABEL_RE.match(p)
        if m:
            kb_idx = i
            kb_inline = m.group(2).strip()
            break

    if def_idx is not None:
        def_end = kb_idx if kb_idx is not None else len(paras)
        def_parts = ([def_inline] if def_inline else []) + paras[def_idx + 1 : def_end]
        definition = " ".join(def_parts)
    else:
        # no explicit "Definition" label: definition text starts right after the name
        def_end = kb_idx if kb_idx is not None else len(paras)
        definition = " ".join(paras[1:def_end])

    if kb_idx is not None:
        key_behaviours = ([kb_inline] if kb_inline else []) + paras[kb_idx + 1 :]
    else:
        key_behaviours = []

    levels = []
    for row in table.rows[2:]:
        level_name_raw, sublevel_raw, indicators_cell = row.cells
        level_name = level_name_raw.text.strip()
        sublevel_text = sublevel_raw.text.strip()  # e.g. "1&2"
        lo, hi = sublevel_text.split("&")
        indicators = [p.text.strip() for p in indicators_cell.paragraphs if p.text.strip()]
        levels.append(
            {
                "subLevelMin": int(lo),
                "subLevelMax": int(hi),
                "levelName": level_name,
                "behaviourIndicators": indicators,
            }
        )
    assert len(levels) == 5, (name, len(levels))
    for lvl, code in zip(levels, LEVEL_CODES):
        lvl["level"] = code

    normalized_name = normalize(name)
    aliased_core_name = normalize(alias_map.get(name, ""))
    is_core = normalized_name in core_names_normalized or aliased_core_name in core_names_normalized

    competencies.append(
        {
            "sno": entry["sno"],
            "cluster": entry["cluster"],
            "subCluster": entry["subCluster"],
            "name": name,
            "isCore": is_core,
            "definition": definition,
            "keyBehaviours": key_behaviours,
            "levels": levels,
        }
    )

out = {
    "proficiencyScale": scale,
    "coreCompetencyNames": [e["name"] for e in core_entries],
    "competencies": competencies,
}

with open(OUT, "w") as f:
    json.dump(out, f, indent=2, ensure_ascii=False)

print("wrote", OUT)
print("competencies:", len(competencies), "core flagged:", sum(1 for c in competencies if c["isCore"]))
for c in competencies:
    if c["isCore"]:
        print(" core:", c["name"])
